from pathlib import Path
import pytest
from docx import Document
from modulos.docx_personalizado import generar_pdf_personalizado


def _crear_plantilla(tmp_path: Path, placeholders: list[str] | None = None) -> Path:
    """Crea una plantilla DOCX mínima con los placeholders indicados."""
    if placeholders is None:
        placeholders = ["NOMBRE", "PERFIL"]
    doc = Document()
    for ph in placeholders:
        doc.add_paragraph("{{" + ph + "}}")
    p = tmp_path / "plantilla.docx"
    doc.save(p)
    return p


def test_generar_pdf_personalizado_happy_path(tmp_path):
    plantilla = _crear_plantilla(tmp_path, ["NOMBRE", "PERFIL"])
    salida_pdf = tmp_path / "salida.pdf"

    generar_pdf_personalizado(
        plantilla,
        valores={
            "NOMBRE": "Daniel Villalba",
            "PERFIL": "Docente con experiencia en TIC.",
        },
        salida_pdf=salida_pdf,
    )

    assert salida_pdf.exists(), "El PDF de salida debe existir"
    assert salida_pdf.stat().st_size > 0, "El PDF de salida no debe estar vacío"

    # El DOCX intermedio NO debe quedar
    tmp_docx = salida_pdf.with_suffix(".__tmp__.docx")
    assert not tmp_docx.exists(), "El DOCX intermedio debe eliminarse tras la conversión"


def test_generar_pdf_personalizado_falla_si_falta_placeholder(tmp_path):
    plantilla = _crear_plantilla(tmp_path, ["X"])
    salida_pdf = tmp_path / "salida.pdf"

    with pytest.raises(ValueError, match="placeholder"):
        generar_pdf_personalizado(
            plantilla,
            valores={},
            salida_pdf=salida_pdf,
        )

    # Si llegó a crearse, el DOCX intermedio debe haberse limpiado
    tmp_docx = salida_pdf.with_suffix(".__tmp__.docx")
    assert not tmp_docx.exists(), "El DOCX intermedio debe limpiarse incluso si falla"


def test_generar_pdf_personalizado_crea_dir_salida(tmp_path):
    plantilla = _crear_plantilla(tmp_path, ["NOMBRE"])
    salida_pdf = tmp_path / "subdir_inexistente" / "anidado" / "salida.pdf"

    assert not salida_pdf.parent.exists(), "Precondición: el directorio destino no existe"

    generar_pdf_personalizado(
        plantilla,
        valores={"NOMBRE": "Daniel"},
        salida_pdf=salida_pdf,
    )

    assert salida_pdf.parent.exists(), "El directorio de salida debe haberse creado"
    assert salida_pdf.exists(), "El PDF debe quedar en la ubicación solicitada"
    assert salida_pdf.stat().st_size > 0
