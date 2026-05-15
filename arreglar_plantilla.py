"""One-shot: arregla los placeholders {{EXP_N_BULLETS}} en plantilla_base.docx.

Cada placeholder de bullets debe tener estilo `List Paragraph` y alineación
default, para que al rellenar quede como los demás bullets de su grupo. Los
párrafos vacíos adyacentes ya tienen el estilo correcto — solo arreglamos el
del placeholder mismo.
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

RUTA = Path(__file__).parent / "data" / "plantilla_base.docx"
RE_BULLETS = re.compile(r"\{\{EXP_\d+_BULLETS\}\}")

doc = Document(str(RUTA))
parrafos = list(doc.paragraphs)
arreglados = 0

for i, p in enumerate(parrafos):
    if not RE_BULLETS.search(p.text):
        continue
    # Buscar el primer párrafo vacío adyacente con estilo List Paragraph
    estilo_objetivo = None
    for j in range(i + 1, min(i + 6, len(parrafos))):
        if parrafos[j].style.name == "List Paragraph" and not parrafos[j].text.strip():
            estilo_objetivo = parrafos[j].style
            break
    if estilo_objetivo is None:
        print(f"[!] {p.text!r}: no se encontró un vecino List Paragraph; se omite")
        continue
    cambios = []
    if p.style.name != estilo_objetivo.name:
        p.style = estilo_objetivo
        cambios.append(f"style={estilo_objetivo.name}")
    if p.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
        p.alignment = None
        cambios.append("align=default")
    if cambios:
        print(f"[OK] {p.text!r}: {', '.join(cambios)}")
        arreglados += 1

doc.save(str(RUTA))
print(f"\nArreglados: {arreglados}")
